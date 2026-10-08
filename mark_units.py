"""指定した台番のセルを黄色に色付けした Excel を作成する

使用例:
    # pision からダウンロード済みの Excel に色付け
    python mark_units.py --file 上大岡.xlsx --units 1,3,5,7 -o 上大岡_当たり.xlsx

    # pision API からデータを取得して Excel を作成し、色付け（PISION_API_KEY が必要）
    python mark_units.py --hall プラザ上大岡 --date 2026-10-06 --units 1001,1003,1005,1007

    # 台番セルではなく行全体を色付け
    python mark_units.py --file 上大岡.xlsx --units 1,3,5 --row
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

YELLOW = PatternFill(fill_type='solid', start_color='FFFFFF00', end_color='FFFFFF00')
DAI_HEADERS = {'台番', '台番号', '台No', '台NO', '台No.', '台NO.', '台'}
_MAX_HEADER_SCAN = 50


def _norm_unit(v) -> str | None:
    """台番を比較用に正規化（'001' / 1 / 1.0 → '1'）"""
    if v is None:
        return None
    s = str(v).strip().translate(str.maketrans('０１２３４５６７８９', '0123456789'))
    if s.endswith('.0'):
        s = s[:-2]
    if s.isdigit():
        return str(int(s))
    return s or None


def parse_units(text: str) -> list[str]:
    units = [_norm_unit(u) for u in text.replace('、', ',').replace(' ', ',').split(',')]
    return [u for u in units if u]


def _find_dai_column(ws) -> tuple[int, int] | None:
    """(header_row, col) を 1-based で返す"""
    for row in ws.iter_rows(max_row=_MAX_HEADER_SCAN):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in DAI_HEADERS:
                return c.row, c.column
    return None


def mark_workbook(wb, units: list[str], col_letter: str | None = None,
                  whole_row: bool = False) -> tuple[set[str], list[str]]:
    """全シートで対象台番を色付け。(見つかった台番, ログ) を返す"""
    targets = set(units)
    found: set[str] = set()
    log: list[str] = []
    for ws in wb.worksheets:
        if col_letter:
            header_row, col = 0, column_index_from_string(col_letter.upper())
        else:
            hit = _find_dai_column(ws)
            if not hit:
                log.append(f'[{ws.title}] 台番列が見つからないためスキップ')
                continue
            header_row, col = hit
        n = 0
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            u = _norm_unit(cell.value)
            if u not in targets:
                continue
            cells = ws[r] if whole_row else (cell,)
            for c in cells:
                c.fill = YELLOW
            found.add(u)
            n += 1
        log.append(f'[{ws.title}] {n} セルを色付け（台番列: {get_column_letter(col)}）')
    return found, log


def build_workbook_from_api(hall_name: str, target_date: str):
    """pision API からホール別結果を取得し Excel を組み立てる"""
    import pision_api as papi

    halls = [h for h in papi.get_halls() if hall_name in h['name']]
    if not halls:
        raise ValueError(f'ホールが見つかりません: {hall_name}')
    if len(halls) > 1:
        names = ', '.join(f"{h['name']}(ID:{h['id']})" for h in halls)
        raise ValueError(f'ホール候補が複数あります。名前を絞り込んでください: {names}')
    hall = halls[0]
    result = papi.get_hall_results(hall['id'], target_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = target_date
    ws.append(['台番', '機種名', '差枚', 'G数', 'BB', 'RB', 'ART'])
    for c in ws[1]:
        c.font = Font(bold=True)
    details = result.get('details', [])
    for d in details:
        model = d.get('model') or {}
        # 台番は unitId（displayName には機種名が入る）
        ws.append([d.get('unitId'), model.get('name'), d.get('diff'), d.get('games'),
                   d.get('bb'), d.get('rb'), d.get('art')])
    ws.column_dimensions['B'].width = 40
    ids = [d['unitId'] for d in details if isinstance(d.get('unitId'), int)]
    if ids:
        print(f"{hall['name']} {target_date}: {len(ids)} 台（台番 {min(ids)}〜{max(ids)}）")
    return wb, hall['name']


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description='指定台番のセルを黄色に色付け')
    src = p.add_mutually_exclusive_group(required=True)
    src.add_argument('--file', help='色付けする Excel ファイル (.xlsx)')
    src.add_argument('--hall', help='pision から取得するホール名（部分一致）')
    p.add_argument('--date', help='--hall 使用時の対象日 (YYYY-MM-DD)')
    p.add_argument('--units', required=True, help='台番（カンマ区切り）')
    p.add_argument('--col', help='台番列を列記号で指定（例: A）。省略時はヘッダーから自動検出')
    p.add_argument('--row', action='store_true', help='台番セルではなく行全体を色付け')
    p.add_argument('-o', '--output', help='出力ファイル名')
    args = p.parse_args(argv)

    units = parse_units(args.units)
    if args.file:
        wb = openpyxl.load_workbook(args.file)
        out = args.output or str(Path(args.file).with_name(Path(args.file).stem + '_当たり.xlsx'))
    else:
        if not args.date:
            p.error('--hall を使う場合は --date が必要です')
        wb, hall_name = build_workbook_from_api(args.hall, args.date)
        out = args.output or f'{hall_name}_{args.date}_当たり.xlsx'

    found, log = mark_workbook(wb, units, args.col, args.row)
    for line in log:
        print(line)
    missing = [u for u in units if u not in found]
    print(f'対象 {len(units)} 台中 {len(units) - len(missing)} 台を色付け')
    if missing:
        print(f'見つからなかった台番: {", ".join(missing)}')
    wb.save(out)
    print(f'保存: {out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
